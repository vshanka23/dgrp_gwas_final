#!/usr/bin/env python3
import pandas as pd
import argparse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side

def main():
    parser = argparse.ArgumentParser(description="Flatten PLINK .qassoc.means blocks into one row per SNP and export as styled Excel")
    parser.add_argument("--input", required=True, help="Input .qassoc.means file")
    parser.add_argument("--output", required=True, help="Output Excel file (.xlsx)")
    args = parser.parse_args()

    # Read file
    df = pd.read_csv(args.input, delim_whitespace=True)

    # Unique SNPs
    snps = df["SNP"].unique()
    out = pd.DataFrame(index=snps)

    for snp in snps:
        out.loc[snp, "G11"] = df.loc[(df["SNP"] == snp) & (df["VALUE"] == "GENO"), "G11"].values[0]
        out.loc[snp, "G22"] = df.loc[(df["SNP"] == snp) & (df["VALUE"] == "GENO"), "G22"].values[0]
        out.loc[snp, "G11_Count"] = float(df.loc[(df["SNP"] == snp) & (df["VALUE"] == "COUNTS"), "G11"].values[0])
        out.loc[snp, "G22_Count"] = float(df.loc[(df["SNP"] == snp) & (df["VALUE"] == "COUNTS"), "G22"].values[0])
        out.loc[snp, "G11_Mean"] = float(df.loc[(df["SNP"] == snp) & (df["VALUE"] == "MEAN"), "G11"].values[0])
        out.loc[snp, "G22_Mean"] = float(df.loc[(df["SNP"] == snp) & (df["VALUE"] == "MEAN"), "G22"].values[0])
        out.loc[snp, "G11_SE"] = (
            float(df.loc[(df["SNP"] == snp) & (df["VALUE"] == "SD"), "G11"].values[0]) /
            (out.loc[snp, "G11_Count"] ** 0.5)
        )
        out.loc[snp, "G22_SE"] = (
            float(df.loc[(df["SNP"] == snp) & (df["VALUE"] == "SD"), "G22"].values[0]) /
            (out.loc[snp, "G22_Count"] ** 0.5)
        )

    # Reset index and round SE columns
    out = out.reset_index().rename(columns={"index": "SNP"})
    out["G11_SE"] = out["G11_SE"].round(3)
    out["G22_SE"] = out["G22_SE"].round(3)

    # Export to Excel
    out.to_excel(args.output, index=False)

    # Apply formatting with openpyxl
    wb = load_workbook(args.output)
    ws = wb.active

    # Define styles
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    snp_font = Font(name="Arial", size=11, bold=True, color="000000")
    normal_font = Font(name="Arial", size=11, color="000000")
    header_fill = PatternFill("solid", fgColor="000000")
    border = Border(
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000"),
    )

    # Style header row
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    # Style data rows
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = normal_font
            cell.border = border
        # SNP column bold
        row[0].font = snp_font

    wb.save(args.output)

if __name__ == "__main__":
    main()
